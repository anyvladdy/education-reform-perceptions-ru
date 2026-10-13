"""Небольшие общие функции: браузер, ожидание, числа, даты и словарь отбора."""

import csv
import json
import os
import re
import time
from datetime import datetime, timedelta, timezone
from pathlib import Path

import pandas as pd
from bs4 import BeautifulSoup
from selenium import webdriver
from selenium.common.exceptions import TimeoutException, WebDriverException, InvalidSessionIdException
from selenium.webdriver.chrome.service import Service
from selenium.webdriver.common.by import By
from selenium.webdriver.support import expected_conditions as EC
from selenium.webdriver.support.ui import WebDriverWait

MOSCOW = timezone(timedelta(hours=3))
ERROR_DIR = Path('data')
TERMS = [
    'колледж', 'техникум', 'СПО', 'ПТУ', 'высшее образован', 'вуз', 'университет',
    'Профессионалитет', 'рабочи', 'професси', 'профориентац', 'поступ', 'абитуриент',
    '10 класс', '10-й класс', 'десятый класс', '9 класс', '9-го класса', 'ОГЭ', 'ЕГЭ',
    'родител', 'диплом', 'престиж', 'трудоустрой', 'зарплат', 'бюджет', 'платное',
    'не берут', 'не приняли', 'отказ', 'передум', 'не пройду', 'не прош', 'попроще',
    'перевод', 'перевестись', 'WorldSkills', 'EuroSkills', 'Молодые профессионалы',
]


def create_driver():
    """На другом компьютере Selenium сам подберёт драйвер установленного Chrome."""
    options = webdriver.ChromeOptions()
    options.add_argument('--headless=new')
    options.add_argument('--window-size=1440,1000')
    options.page_load_strategy = 'eager'
    options.add_experimental_option('prefs', {'download_restrictions': 3})
    service = Service(os.environ['CHROMEDRIVER']) if os.environ.get('CHROMEDRIVER') else Service()
    driver = webdriver.Chrome(service=service, options=options)
    driver.set_page_load_timeout(45)
    return driver


def load_page(driver, url, selector):
    """Только обычная навигация Selenium; два повтора временного сбоя загрузки."""
    for attempt in range(3):
        try:
            driver.get(url)
            title = driver.title.lower()
            visible_start = driver.find_element(By.TAG_NAME,'body').text[:1000].lower()
            if any(x in title for x in ['privacy error', 'just a moment', 'access denied', 'проверка браузера']):
                raise PermissionError('Сайт требует проверки доступа или показывает ошибку сертификата')
            if any(x in visible_start for x in ['проверяем, что вы не робот','подтвердите, что вы человек','verify you are human','too many requests']):
                raise PermissionError('Сайт запросил проверку доступа или ограничил частоту запросов')
            WebDriverWait(driver, 20).until(EC.presence_of_element_located((By.CSS_SELECTOR, selector)))
            if 'pikabu.ru/story/' in url:
                try:
                    WebDriverWait(driver, 10).until(EC.presence_of_element_located((By.CSS_SELECTOR, '.comments')))
                except TimeoutException:
                    note_error('pikabu', url, 'Блок комментариев не появился за 10 секунд; сохранён доступный DOM')
            return BeautifulSoup(driver.page_source, 'html.parser')
        except (TimeoutException, WebDriverException):
            try:
                title = driver.title.lower()
            except WebDriverException:
                raise InvalidSessionIdException('Браузер недоступен: повторите запуск; готовые страницы сохранены')
            if any(x in title for x in ['privacy error', 'just a moment', 'access denied', 'captcha']):
                raise PermissionError('Остановка на проверке доступа; обход защиты не выполняется')
            if attempt == 2:
                raise
            time.sleep(2 * (attempt + 1))


def parse_count(raw):
    """1.54K → 1540; округлённое обозначение сохраняется отдельным полем."""
    if raw is None or not str(raw).strip():
        return None
    text = str(raw).replace('\u2212', '-').replace('\xa0', '').replace(' ', '').strip()
    match = re.fullmatch(r'([+-]?\d+(?:[.,]\d+)?)([kкmм])?', text, re.I)
    if not match:
        return None
    value = float(match[1].replace(',', '.'))
    multiplier = {'k': 1000, 'к': 1000, 'm': 1000000, 'м': 1000000}.get((match[2] or '').lower(), 1)
    return round(value * multiplier)


def within_period(value, start, end):
    """Верхняя граница включает весь день по московскому времени."""
    if not value:
        return False
    date = datetime.fromisoformat(str(value).replace('Z', '+00:00'))
    if date.tzinfo:
        date = date.astimezone(MOSCOW)
    return start <= date.date().isoformat() <= end


def keyword_hits(text, terms):
    normalized = str(text).lower().replace('ё', 'е')
    hits = []
    for term in terms:
        word = term.lower().replace('ё', 'е')
        if word == 'вуз':
            pattern = r'(?<!\w)вуз\w*'
        else:
            pattern = r'(?<!\w)' + re.escape(word) + r'(?!\w)' if word in {'спо','во','пту','егэ','огэ'} else re.escape(word)
        if re.search(pattern, normalized):
            hits.append(term)
    return hits


def note_error(source, url, reason):
    ERROR_DIR.mkdir(parents=True, exist_ok=True)
    path = ERROR_DIR / f'{source}_errors.csv'
    exists = path.exists()
    with path.open('a', encoding='utf-8-sig', newline='') as stream:
        writer = csv.writer(stream)
        if not exists:
            writer.writerow(['collected_at', 'source', 'url', 'reason'])
        writer.writerow([datetime.now(MOSCOW).isoformat(timespec='seconds'), source, url, reason])


def save_excel(tables, path):
    """Три базы на трёх листах, сохранение pandas. Полный текст остаётся в CSV."""
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(path, engine='openpyxl') as writer:
        for name, frame in tables.items():
            frame = frame.copy()
            if 'date_local' in frame:
                frame['date_local'] = pd.to_datetime(frame['date_local'])
            for column in list(frame.select_dtypes(include=['object', 'string'])):
                # Удаляем только недопустимые управляющие символы Excel, не слова.
                frame[column] = frame[column].map(lambda x: re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f]', '', x) if isinstance(x,str) else x)
                maximum = frame[column].map(lambda x: len(x) if isinstance(x,str) else 0).max()
                if pd.notna(maximum) and maximum > 32767:
                    for part in range(1,(int(maximum)-1)//32000+1):
                        frame[f'{column}_part_{part+1}'] = frame[column].map(lambda x: x[part*32000:(part+1)*32000] if isinstance(x,str) else x)
                    frame[column] = frame[column].map(lambda x:x[:32000] if isinstance(x,str) else x)
            frame.to_excel(writer, sheet_name=name, index=False)
            sheet = writer.sheets[name]
            sheet.freeze_panes = 'C2'
            sheet.auto_filter.ref = sheet.dimensions
            from openpyxl.styles import Alignment, Font, PatternFill
            sheet.sheet_view.showGridLines = False
            sheet.row_dimensions[1].height = 32
            for cell in sheet[1]:
                cell.fill = PatternFill('solid',fgColor='17324D')
                cell.font = Font(color='FFFFFF',bold=True)
                cell.alignment = Alignment(wrap_text=True,vertical='center')
            for row in sheet.iter_rows(min_row=2):
                sheet.row_dimensions[row[0].row].height = 48
                for cell in row:
                    if isinstance(cell.value,str):
                        cell.data_type = 's'  # Текст сайта, начинающийся с '=', не формула.
                    cell.alignment = Alignment(vertical='top',wrap_text=True)
                    if frame.columns[cell.column-1] == 'date_local':
                        cell.number_format = 'dd.mm.yyyy'
                    if cell.row % 2 == 0:
                        cell.fill = PatternFill('solid',fgColor='F1F5F9')
            for col in sheet.columns:
                label = str(col[0].value)
                sheet.column_dimensions[col[0].column_letter].width = 75 if label in ['text','title'] else (42 if label=='url' else min(35,max(15,len(label)+2)))
